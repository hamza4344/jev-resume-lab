"""Realistic independent input format smoke tests for the upload endpoints."""
import io
from pathlib import Path
from fastapi.testclient import TestClient
from docx import Document
from pypdf import PdfWriter
from app.main import app, parse_resume
from app.structured_parser import parse_structure

client=TestClient(app)


def test_nonstandard_section_headings():
    txt='''CORE CAPABILITIES
Python SQL Airflow
CAREER JOURNEY
Data Engineer — Amber Analytics — 2022 to Present
- Built scheduled ETL jobs and monitoring dashboards.
SELECTED WORK
Sales Pipeline Migration (2025)
Goal: Reduce errors.
Work: Implemented Python validations.
Result: Reduced reporting turnaround by 50%.
'''
    p=parse_structure(txt)
    assert len(p.experience)==1
    assert p.experience[0].organization=='Amber Analytics'
    assert len(p.projects)==1
    assert p.projects[0].title=='Sales Pipeline Migration (2025)'
    assert p.projects[0].responsibilities[0].text=='Implemented Python validations.'
    assert p.projects[0].outcomes[0].kind=='outcome'


def test_missing_job_title_is_null_and_warned():
    p=parse_structure('''HISTORY
Consultant | Freelance | 2023 - Current
- Reviewed client processes.
Sample Co. | 2022 - 2023
- Maintained process documentation.
''')
    assert len(p.experience)==2
    assert p.experience[1].title is None
    assert p.experience[1].organization=='Sample Co.'
    assert any('no confidently identified title' in w for w in p.warnings)


def test_project_titles_with_acronyms_and_years():
    p=parse_structure('''PROJECTS
Basic FAQ Search Prototype | 2025
Objective: Build search.
- Created a keyword search prototype.
''')
    assert len(p.projects)==1
    assert p.projects[0].title=='Basic FAQ Search Prototype | 2025'


def test_api_parse_txt_and_review_offline(monkeypatch):
    monkeypatch.delenv('BEATAPI_API_KEY',raising=False)
    text='''SUMMARY
Python backend engineer building microservices.
SKILLS
Python, FastAPI, Docker
EXPERIENCE
Backend Engineer | Sample | 2022 - Present
- Built REST APIs using Python and FastAPI.
- Reduced API latency by 22%.
'''
    payload={'file':('candidate.txt',text.encode('utf-8'),'text/plain')}
    parsed=client.post('/api/parse', files=payload)
    assert parsed.status_code==200, parsed.text
    assert parsed.json()['structured_resume']['experience'][0]['title']=='Backend Engineer'
    reviewed=client.post('/api/review',files=payload,
          data={'target_role':'Python Backend Engineer','job_description':'Python FastAPI Docker REST APIs SQL'})
    assert reviewed.status_code==200, reviewed.text
    result=reviewed.json()
    assert result['model']=='local-heuristic'
    assert result['structured_resume']['experience'][0]['responsibilities']
    assert len(result['bullet_reviews'])==2
    assert result['matching_skills']
    assert result['missing_required_skills']==['SQL'] or 'SQL' in result['missing_required_skills']


def test_docx_file_extraction_and_api():
    doc=Document()
    doc.add_paragraph('EXPERIENCE')
    doc.add_paragraph('Marketing Executive | Sample Company | Jan 2021 - Present')
    doc.add_paragraph('- Increased email click-through rate by 25%.')
    content=io.BytesIO(); doc.save(content)
    parsed=client.post('/api/parse',files={'file':('resume.docx',content.getvalue(),
                                                'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert parsed.status_code==200,parsed.text
    assert parsed.json()['structured_resume']['experience'][0]['title']=='Marketing Executive'


def test_scanned_pdf_without_ocr_fails_clear_error():
    writer=PdfWriter(); writer.add_blank_page(width=612,height=792)
    buf=io.BytesIO(); writer.write(buf)
    response=client.post('/api/parse',files={'file':('scan.pdf',buf.getvalue(),'application/pdf')})
    assert response.status_code==422
    assert 'No extractable text' in response.text


def test_unsupported_extension_and_oversized_file():
    resp=client.post('/api/parse',files={'file':('resume.zip',b'bad','application/zip')})
    assert resp.status_code==415
    resp=client.post('/api/parse',files={'file':('resume.txt',b'A'*(5*1024*1024+1),'text/plain')})
    assert resp.status_code==413


def test_parse_endpoint_requires_app_api_key_if_configured(monkeypatch):
    monkeypatch.setenv('APP_API_KEY','local-test-secret')
    payload={'file':('resume.txt',b'EXPERIENCE\n- Built a Python service.','text/plain')}
    assert client.post('/api/parse',files=payload).status_code==401
    assert client.post('/api/parse',files=payload,headers={'x-api-key':'local-test-secret'}).status_code==200
